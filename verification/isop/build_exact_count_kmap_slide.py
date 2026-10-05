#!/usr/bin/env python3
"""Generate a one-slide PowerPoint explaining intrinsic SOP explosion."""

from math import comb
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("exact_count_kmap_slide.pptx")

FONT_JP = "Yu Gothic"
FONT_LATIN = "Aptos"

NAVY = RGBColor(0x13, 0x2A, 0x43)
BLUE = RGBColor(0x1F, 0x5A, 0x94)
TEAL = RGBColor(0x0F, 0x8B, 0x8D)
ORANGE = RGBColor(0xF2, 0x83, 0x22)
ORANGE_PALE = RGBColor(0xFF, 0xED, 0xDB)
RED = RGBColor(0xCF, 0x3F, 0x43)
GREEN = RGBColor(0x1F, 0x8A, 0x5B)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x68, 0x72, 0x80)
LINE = RGBColor(0xD9, 0xE0, 0xE8)
PALE = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_east_asian_font(run, typeface=FONT_JP):
    """Set both latin and East Asian fonts so Japanese renders predictably."""
    run.font.name = FONT_LATIN
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        old = rpr.find(tag, rpr.nsmap)
        if old is not None:
            rpr.remove(old)
        node = OxmlElement(tag)
        node.set("typeface", typeface if tag == "a:ea" else FONT_LATIN)
        rpr.append(node)


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
             margin=0.0, font=FONT_JP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = 1.0
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    set_east_asian_font(r, font)
    return box


def add_round_rect(slide, x, y, w, h, fill=WHITE, line=LINE,
                   radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_width=1.0):
    shape = slide.shapes.add_shape(radius_shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_pill(slide, x, y, w, h, text, fill, color=WHITE, size=11):
    add_round_rect(slide, x, y, w, h, fill=fill, line=fill)
    add_text(slide, x, y + 0.005, w, h - 0.01, text, size=size, color=color,
             bold=True, align=PP_ALIGN.CENTER)


def set_cell_text(cell, text, size, color, bold=False):
    cell.text = ""
    tf = cell.text_frame
    tf.clear()
    tf.margin_left = Pt(1)
    tf.margin_right = Pt(1)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    set_east_asian_font(r)


def set_cell_border(cell, color="D9E0E8", width="12700"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tc_pr.find(edge, tc_pr.nsmap)
        if old is not None:
            tc_pr.remove(old)
        ln = OxmlElement(edge)
        ln.set("w", width)
        solid = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", color)
        solid.append(srgb)
        ln.append(solid)
        tc_pr.append(ln)


def add_kmap(slide):
    # Card
    add_round_rect(slide, 0.55, 1.52, 5.30, 4.78)
    add_text(slide, 0.88, 1.75, 4.65, 0.33,
             "カルノー図：1をまとめられない", size=18, color=NAVY, bold=True)
    add_text(slide, 0.88, 2.09, 4.65, 0.27,
             "E₄,₂ = 1  ⇔  4入力のうち、ちょうど2個が1", size=12, color=MUTED)

    row_labels = ["00", "01", "11", "10"]
    col_labels = ["00", "01", "11", "10"]
    vals = [
        [0, 0, 1, 0],
        [0, 1, 0, 1],
        [1, 0, 0, 0],
        [0, 1, 0, 1],
    ]

    x, y, w, h = 1.15, 2.60, 4.10, 2.42
    table = slide.shapes.add_table(5, 5, Inches(x), Inches(y), Inches(w), Inches(h)).table
    table.first_row = False
    table.first_col = False
    table.horz_banding = False
    table.vert_banding = False
    for col in table.columns:
        col.width = Inches(w / 5)
    for row in table.rows:
        row.height = Inches(h / 5)

    for i in range(5):
        for j in range(5):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PALE
            set_cell_border(cell)

    table.cell(0, 0).fill.fore_color.rgb = NAVY
    set_cell_text(table.cell(0, 0), "AB＼CD", 10, WHITE, True)
    for j, label in enumerate(col_labels, 1):
        table.cell(0, j).fill.fore_color.rgb = NAVY
        set_cell_text(table.cell(0, j), label, 12, WHITE, True)
    for i, label in enumerate(row_labels, 1):
        table.cell(i, 0).fill.fore_color.rgb = NAVY
        set_cell_text(table.cell(i, 0), label, 12, WHITE, True)

    for i in range(4):
        for j in range(4):
            cell = table.cell(i + 1, j + 1)
            if vals[i][j]:
                cell.fill.fore_color.rgb = ORANGE
                set_cell_text(cell, "1", 19, WHITE, True)
            else:
                cell.fill.fore_color.rgb = PALE
                set_cell_text(cell, "0", 15, RGBColor(0xA2, 0xAC, 0xB8), False)

    add_pill(slide, 1.19, 5.26, 3.98, 0.40,
             "隣接する1がゼロ → 2セル以上をグループ化できない", RED, size=11)
    add_text(slide, 0.91, 5.78, 4.60, 0.29,
             "1ビット反転すると、1の個数は 2 → 1 または 3", size=11, color=MUTED,
             align=PP_ALIGN.CENTER)


def add_example_and_growth(slide):
    add_round_rect(slide, 6.05, 1.52, 6.73, 4.78)
    add_text(slide, 6.38, 1.75, 6.07, 0.33,
             "なぜ X（ドントケア）を入れられないのか", size=18, color=NAVY, bold=True)

    # Example flow
    add_round_rect(slide, 6.43, 2.23, 1.35, 0.63, fill=ORANGE_PALE, line=ORANGE)
    add_text(slide, 6.43, 2.23, 1.35, 0.35, "0011", size=20, color=ORANGE, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, 6.43, 2.55, 1.35, 0.22, "Σ=2  ✓", size=10, color=GREEN, bold=True,
             align=PP_ALIGN.CENTER)

    add_text(slide, 7.90, 2.33, 0.55, 0.33, "→", size=25, color=BLUE, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, 7.76, 2.67, 0.84, 0.19, "1bitをXへ", size=9, color=MUTED,
             align=PP_ALIGN.CENTER)

    add_round_rect(slide, 8.58, 2.23, 1.35, 0.63, fill=PALE, line=BLUE)
    add_text(slide, 8.58, 2.23, 1.35, 0.35, "X011", size=20, color=BLUE, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, 8.58, 2.55, 1.35, 0.22, "2入力を含む", size=10, color=MUTED,
             align=PP_ALIGN.CENTER)

    add_text(slide, 10.02, 2.27, 0.42, 0.50, "{", size=35, color=MUTED,
             align=PP_ALIGN.CENTER)
    add_text(slide, 10.42, 2.24, 1.73, 0.25, "0011  Σ=2  ✓", size=12, color=GREEN, bold=True)
    add_text(slide, 10.42, 2.55, 1.73, 0.25, "1011  Σ=3  ✕", size=12, color=RED, bold=True)
    add_text(slide, 6.42, 3.00, 5.88, 0.36,
             "非検出入力を混ぜるため、どの1セルにも X を入れられない", size=13,
             color=RED, bold=True, align=PP_ALIGN.CENTER)

    # Divider
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.40), Inches(3.50), Inches(5.98), Inches(0.012))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()

    add_text(slide, 6.38, 3.72, 3.35, 0.34,
             "必要キューブ数は組合せ数", size=16, color=NAVY, bold=True)
    add_round_rect(slide, 9.74, 3.66, 2.63, 0.52, fill=RGBColor(0xE7, 0xF4, 0xF4), line=TEAL)
    add_text(slide, 9.74, 3.66, 2.63, 0.52,
             "#cubes = C(n, k)", size=19, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    ns = [4, 10, 16, 18, 20]
    ks = [n // 2 for n in ns]
    counts = [comb(n, k) for n, k in zip(ns, ks)]
    labels = ["4入力", "10入力", "16入力", "18入力", "20入力"]

    table = slide.shapes.add_table(2, len(ns), Inches(6.42), Inches(4.40), Inches(5.90), Inches(1.12)).table
    table.first_row = False
    table.first_col = False
    table.horz_banding = False
    for col in table.columns:
        col.width = Inches(5.90 / len(ns))
    for row in table.rows:
        row.height = Inches(0.56)

    for j, (label, count) in enumerate(zip(labels, counts)):
        top = table.cell(0, j)
        top.fill.solid()
        top.fill.fore_color.rgb = NAVY
        set_cell_border(top, color="FFFFFF")
        set_cell_text(top, label, 10, WHITE, True)

        bottom = table.cell(1, j)
        bottom.fill.solid()
        bottom.fill.fore_color.rgb = ORANGE_PALE if j >= 3 else PALE
        set_cell_border(bottom, color="FFFFFF")
        set_cell_text(bottom, f"{count:,}", 13, ORANGE if j >= 3 else INK, True)

    add_text(slide, 6.42, 5.70, 5.92, 0.36,
             "18入力だけで 48,620 本 — 回路規模ではなく、表現形式がボトルネック", size=12,
             color=ORANGE, bold=True, align=PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)

    add_pill(slide, 0.56, 0.31, 1.70, 0.32, "XORを使わない反例", NAVY, size=10)
    add_text(slide, 0.56, 0.68, 12.10, 0.50,
             "小さな回路でも、キューブ表現は指数的に膨らむ", size=28, color=NAVY, bold=True)
    add_text(slide, 0.58, 1.18, 11.85, 0.25,
             "exact-count関数：AND / OR / NOTで構成できるが、平坦なSOPでは1セルずつ列挙が必要", size=13,
             color=MUTED)

    add_kmap(slide)
    add_example_and_growth(slide)

    # Takeaway footer
    add_round_rect(slide, 0.55, 6.52, 12.23, 0.64, fill=NAVY, line=NAVY)
    add_text(slide, 0.82, 6.64, 4.15, 0.35,
             "回路：AND / OR / NOTで多項式サイズ", size=15, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, 5.03, 6.62, 0.55, 0.35, "≠", size=22, color=ORANGE, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, 5.59, 6.64, 6.82, 0.35,
             "平坦SOP：C(n,n/2) ≈ 2ⁿ/√n 本（指数的）", size=15, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
