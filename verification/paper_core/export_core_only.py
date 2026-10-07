#!/usr/bin/env python3
"""Validate finished C runs and export a three-way comparison (no charts)."""
import argparse
import csv
import json
from pathlib import Path
import re
import statistics
from decimal import Decimal
from itertools import combinations
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
MODES = ('xid', 'core_only', 'core_min')


def read_csv(path):
    with path.open() as fp:
        return list(csv.DictReader(fp))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run-dir', type=Path, required=True)
    args = parser.parse_args()
    out = args.run_dir.resolve()
    repeats = int((out / 'repeats.txt').read_text())
    def config(name, default):
        path = out / (name + '.txt')
        return path.read_text().strip() if path.exists() else default
    profile = config('profile', 'small')
    limit = int(config('limit', '0'))
    warmups = int(config('warmups', '1'))
    verification = config('verification', 'on') == 'on'
    references, fixed, measured, validated, faults_by_run = {}, {}, [], [], {}
    for path in sorted(out.glob('*.timing.csv')):
        name = path.name.removesuffix('.timing.csv')
        match = re.fullmatch(r'(.+)_(xid|core_only|core_min)_(warmup|verify|\d+)', name)
        assert match, path
        circuit, mode, label = match.groups()
        stem = out / name
        timing = read_csv(path)[0]
        assert int(timing['exit_code']) == 0, stem
        rows = read_csv(stem.with_suffix('.csv'))
        assert all(r['complete'] in ('0', '1', '') for r in rows), stem
        if limit == 0:
            assert all(r['complete'] in ('1', '') for r in rows), stem
        representative = [r for r in rows if r['complete'] in ('0', '1')]
        fault_map = {(r['net_name'], r['f_type']): r for r in representative}
        assert len(fault_map) == len(representative), (stem, 'duplicate representative')
        assert all(int(r['seeded_cnt']) == 0 for r in representative), stem
        if profile == 'medium':
            assert len(representative) == {'s5378_C': 4551, 's9234_C': 6927}[circuit], stem
            assert all(0 <= int(r['cube_cnt']) <= limit for r in representative), stem
        faults_by_run[circuit, mode, label] = fault_map
        fdp = {(r['net_name'], r['f_type']): r['fdp'] for r in rows}
        if circuit not in references:
            references[circuit] = fdp
        assert fdp.keys() == references[circuit].keys(), (stem, 'physical fault keys differ')
        if limit == 0:
            assert fdp == references[circuit], (stem, 'FDP differs')
        cubes = sum(int(r['cube_cnt']) for r in representative)
        log = stem.with_suffix('.log').read_text()
        stderr = stem.with_suffix('.stderr').read_text()
        def setting(key, value):
            assert re.search(r'//\s+' + re.escape(key) + r'\s+: ' + re.escape(value) + r'\s*$', log, re.M), (stem, key)
        setting("Don't-care Method", 'XID' if mode == 'xid' else 'CORE')
        setting('Dominance Cube Reuse', 'off')
        setting('CORE Extra Verification', 'on' if label == 'verify' else 'off')
        setting('CORE Minimization', 'off' if mode == 'core_only' else 'on')
        setting('CORE Recheck', 'off' if mode == 'core_only' else 'on')
        setting('Low Power', 'off')
        solves, core_care, final_care, input_bits, deletions = [None] * 5
        if mode != 'xid':
            stats = re.search(r'\[PAPER_CORE\] cubes=(\d+) solves=(\d+) care: input=(\d+) core=(\d+) prime=(\d+)', stderr)
            assert stats, stem
            count, solves, input_bits, core_care, final_care = map(int, stats.groups())
            detail = re.search(r'\[PAPER_CORE_MODE\] minimize=(on|off) recheck=(on|off) deletion_queries=(\d+) final_care=(\d+)', stderr)
            assert detail and count == cubes and final_care == int(detail[4]), stem
            deletions = int(detail[3])
            if mode == 'core_only':
                assert core_care == final_care and deletions == 0, stem
                expected = cubes * (2 if label == 'verify' else 1)
            else:
                assert deletions == core_care and final_care <= core_care <= input_bits, stem
                expected = 2 * cubes + core_care
                if label == 'verify':
                    expected += cubes + final_care
            assert solves == expected, (stem, solves, expected)
        if label == 'verify':
            assert '[GT] summary:' in stderr and 'ALL VERIFIED' in stderr, stem
            validated.append({'circuit': circuit, 'mode': mode, 'faults': len(representative), 'cubes': cubes, 'BDD': 'ALL VERIFIED'})
        else:
            assert '[GT]' not in stderr, stem
        def seconds(field):
            m = re.search(r'//\s+' + re.escape(field) + r'\s+: ([0-9.]+) sec', log)
            assert m, (stem, field)
            return float(m[1])
        row = dict(circuit=circuit, mode=mode, run=label, faults=len(representative), cubes=cubes,
                   limit=limit, complete=sum(r['complete'] == '1' for r in representative),
                   incomplete=sum(r['complete'] == '0' for r in representative),
                   cpu_s=float(timing['cpu_s']), wall_s=float(timing['wall_s']),
                   dc_s=seconds("CPU Time (Don't care)"), sat_s=seconds('CPU Time (CaDiCaL)'),
                   bdd_s=seconds('CPU Time (BDD)'), negative_solves=solves,
                   input_bits=input_bits, core_care=core_care, final_care=final_care, deletion_queries=deletions)
        key = circuit, mode
        counts = tuple(row[k] for k in ('faults', 'complete', 'incomplete', 'cubes', 'negative_solves', 'input_bits', 'core_care', 'final_care', 'deletion_queries'))
        # Extra verification changes query count only; compare actual output counts.
        if label != 'verify':
            assert key not in fixed or fixed[key] == counts, (key, 'counts varied')
            fixed[key] = counts
        if label.isdigit():
            measured.append(row)
    assert measured and len(validated) == (len(references) * len(MODES) if verification else 0)
    coverage = []
    details = []
    for circuit in sorted(references):
        for trial in range(1, repeats + 1):
            for lhs, rhs in combinations(MODES, 2):
                a, b = faults_by_run[circuit, lhs, str(trial)], faults_by_run[circuit, rhs, str(trial)]
                assert a.keys() == b.keys(), (circuit, 'representative fault keys differ')
                counts = dict(both_complete=0, lhs_only_complete=0, rhs_only_complete=0, neither_complete=0,
                              partial_lhs_higher=0, partial_rhs_higher=0, partial_equal=0,
                              exact_mismatch=0, missing_faults=0)
                for key in sorted(a):
                    ca, cb = a[key]['complete'] == '1', b[key]['complete'] == '1'
                    category = 'both_complete' if ca and cb else 'lhs_only_complete' if ca else 'rhs_only_complete' if cb else 'neither_complete'
                    counts[category] += 1
                    va, vb = Decimal(a[key]['fdp']), Decimal(b[key]['fdp'])
                    if ca and cb:
                        assert va == vb, (circuit, lhs, rhs, key, 'completed FDP mismatch')
                    else:
                        if ca:
                            assert vb <= va, (circuit, rhs, key, 'partial FDP exceeds completed exact value')
                        if cb:
                            assert va <= vb, (circuit, lhs, key, 'partial FDP exceeds completed exact value')
                        counts['partial_lhs_higher' if va > vb else 'partial_rhs_higher' if va < vb else 'partial_equal'] += 1
                    if profile == 'medium':
                        details.append(dict(circuit=circuit, run=trial, lhs=lhs, rhs=rhs,
                                            net_name=key[0], f_type=key[1], lhs_complete=int(ca), rhs_complete=int(cb),
                                            lhs_cubes=int(a[key]['cube_cnt']), rhs_cubes=int(b[key]['cube_cnt']),
                                            lhs_fdp=str(va), rhs_fdp=str(vb)))
                coverage.append(dict(circuit=circuit, run=trial, lhs=lhs, rhs=rhs, **counts))
    summary = []
    for circuit in sorted(references):
        for mode in MODES:
            rows = [r for r in measured if (r['circuit'], r['mode']) == (circuit, mode)]
            assert len(rows) == repeats and {int(r['run']) for r in rows} == set(range(1, repeats + 1))
            row = {k: rows[0][k] for k in ('circuit', 'mode', 'faults', 'complete', 'incomplete', 'limit', 'cubes', 'negative_solves', 'input_bits', 'core_care', 'final_care', 'deletion_queries')}
            row['repeats'] = repeats
            row['cpu_stat'] = 'single' if repeats == 1 else 'median'
            for field in ('cpu_s', 'wall_s', 'dc_s', 'sat_s', 'bdd_s'):
                row[field] = statistics.median(r[field] for r in rows)
            row['cpu_min_s'], row['cpu_max_s'] = min(r['cpu_s'] for r in rows), max(r['cpu_s'] for r in rows)
            row['core_x_ratio'] = 1 - row['core_care'] / row['input_bits'] if row['input_bits'] else None
            row['final_x_ratio'] = 1 - row['final_care'] / row['input_bits'] if row['input_bits'] else None
            summary.append(row)
    for name, rows in [('measurements', measured), ('summary', summary), ('coverage', coverage), ('fault_comparison', details)]:
        if not rows:
            continue
        with (out / (name + '.csv')).open('w', newline='') as fp:
            writer = csv.DictWriter(fp, fieldnames=list(rows[0]))
            writer.writeheader(); writer.writerows(rows)
    metadata = dict(parent_commit=(out / 'source_parent.txt').read_text().strip(),
                    cpu=int((out / 'cpu.txt').read_text()), repeats=repeats, warmups=warmups, profile=profile,
                    fault_model='SAF', low_power=False, limit=limit, dominance_reuse=False,
                    timed_verification=False, order='rotating three modes',
                    separate_BDD_verification=verification,
                    cpu_metric='wait4 user+system of one C process; startup through exit',
                    binary_hashes=(out / 'binary.sha256').read_text(), run_dir=str(out))
    payload = dict(metadata=metadata, summary=summary, measurements=measured, separate_BDD_verification=validated,
                   coverage_comparison=coverage)
    (out / 'summary.json').write_text(json.dumps(payload, indent=2) + '\n')
    target = ROOT / 'verification/paper_core/results' / ('core_only_medium_comparison.json' if profile == 'medium' else 'core_only_comparison.json')
    target.write_text(json.dumps(payload, indent=2) + '\n')
    # Preserve the small campaign and show both circuit sizes in one workbook.
    campaigns = []
    for name in ('core_only_comparison.json', 'core_only_medium_comparison.json'):
        path = ROOT / 'verification/paper_core/results' / name
        if path.exists():
            campaigns.append(json.loads(path.read_text()))
    all_summary, all_measured, all_validated, all_coverage = [], [], [], []
    for campaign in campaigns:
        for original in campaign['summary']:
            item = dict(original)
            item.setdefault('complete', item['faults'])
            item.setdefault('incomplete', 0)
            item.setdefault('limit', 0)
            item.setdefault('repeats', campaign['metadata']['repeats'])
            item.setdefault('cpu_stat', 'single' if item['repeats'] == 1 else 'median')
            all_summary.append(item)
        for original in campaign['measurements']:
            item = dict(original)
            item.setdefault('complete', item['faults'])
            item.setdefault('incomplete', 0)
            item.setdefault('limit', 0)
            all_measured.append(item)
        all_validated.extend(campaign['separate_BDD_verification'])
        all_coverage.extend(campaign.get('coverage_comparison', []))
    wb = Workbook(); wb.remove(wb.active)
    def sheet(name, fields, rows):
        ws = wb.create_sheet(name); ws.append([label for _, label in fields])
        for row in rows:
            ws.append([row.get(key) for key, _ in fields])
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 23
        for cells in ws.iter_rows(min_row=2):
            for cell, (key, _) in zip(cells, fields):
                if isinstance(cell.value, float):
                    cell.number_format = '0.0000000000E+00' if key.endswith('_fdp') else '0.000000'
    common = [('circuit', '回路'), ('mode', '方式'), ('limit', 'limit（0は無制限）')]
    sheet('実行時間', common + [('repeats', '測定回数'), ('cpu_stat', '時間の集計方法'), ('cpu_s', 'CPU秒'), ('cpu_min_s', 'CPU最小 秒'), ('cpu_max_s', 'CPU最大 秒'), ('wall_s', '実時間 秒')], all_summary)
    sheet('キューブ数', common + [('faults', '代表故障数'), ('cubes', '総キューブ数')], all_summary)
    sheet('故障完了', common + [('faults', '代表故障数'), ('complete', '完了故障数'), ('incomplete', '未完了故障数')], all_summary)
    wide = []
    for circuit in sorted({r['circuit'] for r in all_summary}):
        row = {'circuit': circuit}
        for mode in MODES:
            item = next(r for r in all_summary if (r['circuit'], r['mode']) == (circuit, mode))
            row['limit'], row['repeats'] = item['limit'], item['repeats']
            row[mode + '_cpu_s'] = item['cpu_s']
            row[mode + '_cubes'] = item['cubes']
            row[mode + '_complete'] = item['complete']
        row['core_only_cpu_over_xid'] = row['core_only_cpu_s'] / row['xid_cpu_s']
        row['core_only_cubes_over_xid'] = row['core_only_cubes'] / row['xid_cubes']
        row['core_min_cpu_over_xid'] = row['core_min_cpu_s'] / row['xid_cpu_s']
        wide.append(row)
    sheet('方式別比較', [('circuit', '回路')] + [(key, key) for key in wide[0] if key != 'circuit'], wide)
    sheet('判定回数・X率', common + [(k, k) for k in ('negative_solves', 'deletion_queries', 'input_bits', 'core_care', 'final_care', 'core_x_ratio', 'final_x_ratio')], all_summary)
    sheet('時間内訳', common + [('dc_s', 'ドントケアCPU 秒'), ('sat_s', '生成SAT CPU 秒'), ('bdd_s', 'BDD CPU 秒')], all_summary)
    sheet('測定明細', [(key, key) for key in all_measured[0]], all_measured)
    sheet('別実行BDD検証', [('circuit', '回路'), ('mode', '方式'), ('faults', '代表故障数'), ('cubes', '総キューブ数'), ('BDD', 'BDD検証')], all_validated)
    if all_coverage:
        sheet('方式間の完了比較', [(key, key) for key in all_coverage[0]], all_coverage)
    if details:
        excel_details = [dict(row, lhs_fdp=float(row['lhs_fdp']), rhs_fdp=float(row['rhs_fdp'])) for row in details]
        sheet('中規模故障別比較', [(key, key) for key in details[0]], excel_details)
    ws = wb.create_sheet('条件'); ws.append(['項目', '値'])
    for index, campaign in enumerate(campaigns, 1):
        for key, value in campaign['metadata'].items(): ws.append([f'campaign{index}: {key}', str(value)])
    ws.append(['core_only', 'core_minimize off / core_recheck off / core_verify off; negative solves = cubes'])
    ws.append(['core_min', 'core_minimize on / core_recheck on / core_verify off; negative solves = 2*cubes + core_care'])
    ws.append(['注意', 'core_only の結果は既存の極小化実行の中間値ではなく、別の列挙実行'])
    ws.append(['中規模', 'limit30 / 回数は実行時間シート参照 / warmupなし / 別BDDなし。complete=0のFDPは下界。両方式完了FDPを照合。'])
    excel = ROOT / 'verification/paper_core/core_only_comparison.xlsx'
    wb.save(excel)
    check = load_workbook(excel, data_only=True)
    assert check['キューブ数'].max_row == len(all_summary) + 1
    assert not any(c.data_type == 'f' for ws in check for row in ws for c in row)
    for row in summary:
        print(f'{row["circuit"]:8} {row["mode"]:9} CPU={row["cpu_s"]:.6f}s cubes={row["cubes"]}', flush=True)
    print(f'Excel: {excel}\nRaw results: {out}', flush=True)


if __name__ == '__main__':
    main()
