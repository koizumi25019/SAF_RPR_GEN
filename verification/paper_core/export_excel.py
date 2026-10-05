#!/usr/bin/env python3
"""Export the saved comparison to numeric Excel tables (no charts).
Requires openpyxl. Input snapshots are tracked; original output logs are optional.
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=HERE/'paper_core_comparison.xlsx')
    args = parser.parse_args()
    data = json.loads((HERE/'results/comparison_data.json').read_text())
    simulation = json.loads((HERE/'results/simulation_summary.json').read_text())
    perf = {(r['circuit'],r['method']):r for r in data['performance']}
    names = list(dict.fromkeys(r['circuit'] for r in data['performance']))
    core = {r['circuit']:r for r in data['core_statistics']}
    comparisons = {r['circuit']:r for r in data['coverage_comparison']}
    xid_x = {r['circuit']:r['x_bits']/r['input_bits'] for r in data['xid_x_statistics']}
    small = {(r['circuit'],r['mode']):r for r in data['small_summary']}
    medium = {(r['circuit'],r['method']):r for r in data['raw_runs'] if r['limit']==30}
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = 'XID / paper CORE comparison, SAF'
    workbook.properties.description = 'Numeric tables for graphing. Small: 7-run median, unlimited. Medium: single run, limit30.'
    workbook.properties.creator = 'SAF_RPR_GEN'
    formats = {'text':'General','int':'#,##0','seconds':'0.000000','ratio':'0.00%', 'avg':'0.0000','multiplier':'0.0000'}

    def sheet(title, columns, records):
        ws = workbook.create_sheet(title)
        ws.append([label for label,_ in columns])
        for row in records:
            assert len(row)==len(columns)
            ws.append(row)
        ws.freeze_panes='B2'
        for col,(label,kind) in enumerate(columns,1):
            ws.column_dimensions[get_column_letter(col)].width = min(42,max(18,len(label)*1.7))
            ws.cell(1,col).font=Font(bold=True,color='FFFFFF')
            ws.cell(1,col).fill=PatternFill('solid',fgColor='24476B')
            ws.cell(1,col).alignment=Alignment(wrap_text=True,vertical='center')
            for row in range(2,ws.max_row+1):
                cell=ws.cell(row,col)
                if cell.value is not None and kind!='text':
                    assert isinstance(cell.value,(int,float)), (title,label,cell.value)
                cell.number_format=formats[kind]
        ws.row_dimensions[1].height=36
        table=Table(displayName=f'Data{len(workbook.worksheets)}',ref=f'A1:{get_column_letter(len(columns))}{ws.max_row}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
        return ws

    time_rows,cube_rows,completion_rows,x_rows,query_rows,breakdown_rows=[],[],[],[],[],[]
    for name in names:
        x,p=perf[name,'xid'],perf[name,'core']
        c=core[name]
        n=c['inputs_per_cube'];inputs=c['input_bits'];k=c['care_after_core'];prime=c['care_after_min'];nc=c['cubes']
        sx,sp=small.get((name,'xid'),{}),small.get((name,'core'),{})
        time_rows.append([name,x['limit'],x['repeats'],x['cpu_stat'],x['cpu_s'],p['cpu_s'],p['cpu_s']/x['cpu_s'],
                          sx.get('cpu_min_s'),sx.get('cpu_max_s'),sp.get('cpu_min_s'),sp.get('cpu_max_s')])
        cube_rows.append([name,x['limit'],x['cubes'],p['cubes'],1-p['cubes']/x['cubes'],p['cubes']/x['cubes']])
        comparison=comparisons.get(name,{})
        total=x['representatives']
        completion_rows.append([name,x['limit'],total,x['complete'],p['complete'],x['complete']/total,p['complete']/total,
            total-x['complete'],total-p['complete'],comparison.get('both_complete',total),
            comparison.get('core_only_complete',0),comparison.get('xid_only_complete',0),comparison.get('neither_complete',0)])
        x_rows.append([name,n,xid_x.get(name),1-k/inputs,1-prime/inputs,k/inputs,(k-prime)/k if k else 0])
        query_rows.append([name,n,nc,k/nc,prime/nc,(k-prime)/nc,2+k/nc,inputs,k,prime,c['deletion_trials'],c['negative_solves']])
        tx=small.get((name,'xid'),medium.get((name,'xid')))
        tp=small.get((name,'core'),medium.get((name,'core')))
        breakdown_rows.append([name,x['limit'],tx['dc_s'],tp['dc_s'],tx['sat_s'],tp['sat_s'],tx['bdd_s'],tp['bdd_s']])
    sheet('実行時間',[('回路','text'),('limit（0は無制限）','int'),('測定回数／方式','int'),('統計','text'),
        ('XID CPU秒','seconds'),('CORE CPU秒','seconds'),('時間倍率 CORE/XID','multiplier'),
        ('XID 最小CPU秒','seconds'),('XID 最大CPU秒','seconds'),('CORE 最小CPU秒','seconds'),('CORE 最大CPU秒','seconds')],time_rows)
    sheet('キューブ数',[('回路','text'),('limit（0は無制限）','int'),('XID キューブ数','int'),('CORE キューブ数','int'),
        ('キューブ削減率','ratio'),('キューブ数 CORE/XID','multiplier')],cube_rows)
    sheet('故障完了',[('回路','text'),('limit（0は無制限）','int'),('全代表故障数','int'),('XID 完了数','int'),('CORE 完了数','int'),
        ('XID 完了率','ratio'),('CORE 完了率','ratio'),('XID 未完了数','int'),('CORE 未完了数','int'),
        ('両方式完了','int'),('COREのみ完了','int'),('XIDのみ完了','int'),('両方式未完了','int')],completion_rows)
    sheet('X率',[('回路','text'),('入力ビット数','int'),('XID X率（別の生成系列）','ratio'),('core抽出直後 X率','ratio'),
        ('極小化後 X率','ratio'),('削除判定対象／全入力','ratio'),('判定対象のうちX化','ratio')],x_rows)
    sheet('判定回数',[('回路','text'),('入力ビット数','int'),('生成キューブ数','int'),('coreケア平均／キューブ','avg'),
        ('最終ケア平均／キューブ','avg'),('追加X化平均／キューブ','avg'),('否定側solve平均／キューブ','avg'),
        ('全入力ビット位置数','int'),('coreケア総数','int'),('最終ケア総数','int'),('削除試行総数','int'),('否定側solve総数','int')],query_rows)
    sheet('時間内訳',[('回路','text'),('limit（0は無制限）','int'),('XID DC秒','seconds'),('CORE DC秒','seconds'),
        ('XID 生成SAT秒','seconds'),('CORE 生成SAT秒','seconds'),('XID BDD秒','seconds'),('CORE BDD秒','seconds')],breakdown_rows)
    coverage_rows=[]
    for name,r in comparisons.items():
        coverage_rows.append([name,r['both_complete'],r['core_only_complete'],r['xid_only_complete'],r['neither_complete'],
            r['core_higher_partial'],r['core_lower_partial'],r['equal_partial'],r['exact_mismatch'],r['missing_faults']])
    sheet('中規模FDP比較',[('回路','text'),('両方式完了','int'),('COREのみ完了','int'),('XIDのみ完了','int'),('両方式未完了','int'),
        ('未完了を含むFDP CORE大','int'),('未完了を含むFDP CORE小','int'),('未完了を含むFDP同値','int'),
        ('両方式完了FDP不一致','int'),('故障欠落','int')],coverage_rows)
    raw_rows=[]
    for r in data['raw_runs']:
        raw_rows.append([r['circuit'],r['method'],r['run'],r['limit'],r['cpu_measurement'],r['cpu_s'],r['wall_s'],
            r['reported_cpu_s'],r['dc_s'],r['sat_s'],r['bdd_s'],r['representatives'],r['cubes']])
    sheet('全測定値',[('回路','text'),('方式','text'),('反復','int'),('limit（0は無制限）','int'),('CPU測定方法','text'),
        ('比較用CPU秒','seconds'),('実時間秒','seconds'),('アプリCPU秒','seconds'),('DC秒','seconds'),('生成SAT秒','seconds'),
        ('BDD秒','seconds'),('代表故障数','int'),('キューブ数','int')],raw_rows)
    sim_rows=[]
    for r in simulation['circuits']:
        sim_rows.append([r['circuit'],r['n_inputs'],r['representatives'],r['cubes'],r['core_expanded_patterns'],
            r['prime_expanded_patterns'],r['exhaustive_fault_pattern_pairs'],r['unsound_core_cubes'],r['unsound_prime_cubes'],
            r['exact_cover_mismatches'],'PASS'])
    sheet('シミュレーション',[('回路','text'),('入力ビット数','int'),('代表故障数','int'),('キューブ数','int'),
        ('core直後の展開パターン延べ数','int'),('極小化後の展開パターン延べ数','int'),('全故障×全入力の評価組数','int'),
        ('不正coreキューブ数','int'),('不正最終キューブ数','int'),('完全被覆不一致数','int'),('結果','text')],sim_rows)
    notes=[
        ['対象','縮退故障 SAF。c17a / s27_C / s208_C / s298_C / s5378_C / s9234_C。s13207_C は未実行。'],
        ['日付','2026-10-05〜2026-10-06、Asia/Tokyo。'],
        ['CORE','完全モデルから直接core抽出→core再確認→削除による極小化。XIDは経由しない。'],
        ['比較条件','支配キューブ流用 off、追加core検証 off。低消費電力制約は未実装。'],
        ['limit','0 は無制限。30 は故障ごとに最大30キューブ。秒数の制限ではない。'],
        ['小規模の時間','各方式7回の子プロセス user+system CPU時間中央値。CPU 0に固定。ウォームアップ・検証実行は除外。'],
        ['中規模の時間','各方式1回、アプリclock()の全体CPU時間。CPU固定なし。繰返し測定の統計ではない。'],
        ['全体時間','読込・CNF/oracle構築・生成SAT・DC・BDD・出力を含む。'],
        ['時間内訳','既存ログの0.001秒分解能。小規模は区間ごとの中央値。区間外処理もあり、全体時間への厳密な加算はできない。'],
        ['実時間','小規模はPythonの起動・timeout付きwaitの間隔も含む。CPU秒を比較用に使用。'],
        ['完了','complete=1 は厳密FDP、complete=0 は検出集合の下界。等価故障のエコー行は集計から除外。'],
        ['X率の母数','全生成キューブの入力ビット位置数。故障別の率を均等に平均した値ではない。割合は数値0〜1をExcel百分率表示。'],
        ['core直後','CORE＋極小化の実行中の中間値。極小化を省いた別実験ではない。'],
        ['XID X率','小規模のみ補足測定。中規模は未計測のため空欄（0ではない）。方式間で生成パターンの系列が異なる。'],
        ['削除判定対象','coreに残ったケアビットだけ。残存kビットにつき1回ずつ。否定側solveはキューブ当たり2+k回。'],
        ['判定対象のうちX化','(coreケア数−最終ケア数) / coreケア数。'],
        ['シミュレーション','元Verilogから独立に2値ゲート評価。小規模の全入力をビット並列で列挙し、core直後と極小化後のX展開を全て確認。'],
        ['展開数','キューブ間の重複を含む延べパターン数。ランダムサンプリングではない。'],
        ['検証範囲','小規模4回路の577代表故障。全キューブ和集合＝全検出パターン集合、FDPも一致。中規模の全展開シミュレーションは行っていない。'],
        ['元データ','verification/paper_core/results/comparison_data.json / simulation_summary.json。グラフは作成していない。'],
        ['小規模バイナリSHA256',data['small_binary_sha256']],
        ['中規模バイナリSHA256',data['medium_binary_sha256']],
        ['検証バイナリSHA256',simulation['binary_sha256']],
    ]
    ws=sheet('測定条件',[('項目','text'),('内容','text')],notes)
    ws.column_dimensions['A'].width=30
    ws.column_dimensions['B'].width=110
    for row in range(2,ws.max_row+1):
        ws.cell(row,2).alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[row].height=36
    args.output.parent.mkdir(parents=True,exist_ok=True)
    workbook.save(args.output)
    # Read the saved file back, checking row counts and numeric graph columns.
    saved=load_workbook(args.output,data_only=True)
    assert saved['実行時間'].max_row==7 and saved['全測定値'].max_row==61
    assert saved['シミュレーション'].max_row==5
    assert saved['X率']['C6'].value is None and saved['X率']['C7'].value is None
    for row in saved['X率'].iter_rows(min_row=2,min_col=4,max_col=7):
        assert all(cell.data_type=='n' and 0<=cell.value<=1 for cell in row)
    for ws in saved:
        assert not ws._charts and len(ws.tables)==1
    print(f'Saved and verified {args.output}: {len(saved.sheetnames)} sheets, numeric tables, no charts')


if __name__ == '__main__':
    main()
