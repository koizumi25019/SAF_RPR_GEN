#!/usr/bin/env python3
"""Save low-power TDF measurements as JSON, CSV and numeric Excel tables."""
import argparse
import csv
import json
from pathlib import Path
import statistics

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run-dir', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=HERE/'tdf_power_on_results.xlsx')
    parser.add_argument('--snapshot', type=Path, default=HERE/'results/tdf_power_benchmark.json')
    args = parser.parse_args()
    run_dir = args.run_dir.resolve()
    raw = list(csv.DictReader((run_dir/'measurements.csv').open()))
    planned = (run_dir/'circuits.txt').read_text().split() if (run_dir/'circuits.txt').exists() else list(dict.fromkeys(r['circuit'] for r in raw))
    failures=[]; partial_fault_rows=[]
    for path in sorted(run_dir.glob('*/run*/status.csv')):
        for r in csv.DictReader(path.open()):
            if r['status']=='completed':continue
            for k in ('run','exit_code'):r[k]=int(r[k])
            for k in ('wall_s','cpu_s'):r[k]=float(r[k])
            partial = path.parent/'fdp.csv'
            text = partial.read_text() if partial.exists() else ''
            # SIGTERM may leave the last CSV row incomplete; retain full lines only.
            text = text[:text.rfind('\n')+1]
            stored = list(csv.DictReader(text.splitlines())) if text else []
            primary = [f for f in stored if f.get('complete') in ('0','1')]
            r.update(stored_representatives=len(primary),stored_complete=sum(f['complete']=='1' for f in primary),
                     stored_incomplete=sum(f['complete']=='0' for f in primary),stored_fault_rows=len(stored))
            for f in stored:
                partial_fault_rows.append([r['circuit'],f['net_name'],f['f_type'],int(f['cube_cnt']) if f['cube_cnt'] else None,
                    int(f['complete']) if f['complete'] else None,float(f['fdp'])])
            failures.append(r)
    assert raw or failures, 'No completed or failed measurements'
    for row in raw:
        for key, value in list(row.items()):
            if key != 'circuit':
                row[key] = None if value == '' else (float(value) if key.endswith(('_s','_pct')) else int(value))
    names = list(dict.fromkeys(r['circuit'] for r in raw))
    summary, fault_rows = [], []
    for name in names:
        runs = [r for r in raw if r['circuit']==name]
        expected = 7 if name in ('s27','s208') else 1
        assert [r['run'] for r in runs] == list(range(1,expected+1)), (name,'Missing runs')
        stable = [k for k in runs[0] if k not in ('run','cpu_s','wall_s','dc_s','sat_s','bdd_s')]
        assert all(all(r[k]==runs[0][k] for k in stable) for r in runs), (name,'Non-deterministic statistics')
        row = {k:runs[0][k] for k in stable}
        row.update(repeats=len(runs),statistic='median' if len(runs)>1 else 'single',
                   cpu_min_s=min(r['cpu_s'] for r in runs),cpu_max_s=max(r['cpu_s'] for r in runs))
        for key in ('cpu_s','wall_s','dc_s','sat_s','bdd_s'):
            row[key] = statistics.median(r[key] for r in runs)
        assert row['complete']+row['incomplete']==row['representatives']
        assert row['zero_fdp']+row['positive_fdp']==row['representatives']
        reference = None
        for r in runs:
            csv_path = run_dir/name/f'run{r["run"]}'/'fdp.csv'
            all_faults = list(csv.DictReader(csv_path.open()))
            signature = [(f['net_name'],f['f_type'],f['cube_cnt'],f['complete'],f['fdp']) for f in all_faults]
            if reference is None: reference = signature
            else: assert reference == signature, (name,'Fault results vary across runs')
        # Small circuits must also agree with the independent exhaustive result.
        verified = ROOT/f'output/tdf_power_checks_all_faults/{name}_p20/fdp.csv'
        if name in ('s27','s208') and verified.exists():
            previous = list(csv.DictReader(verified.open()))
            assert {(r['net_name'],r['f_type']):r['fdp'] for r in previous} == {
                (r['net_name'],r['f_type']):r['fdp'] for r in all_faults}, name
        for f in all_faults:
            fault_rows.append([name,f['net_name'],f['f_type'],int(f['cube_cnt']) if f['cube_cnt'] else None,
                               int(f['complete']) if f['complete'] else None,float(f['fdp']),
                               'representative' if f['complete'] else 'equivalent_echo'])
        summary.append(row)
    source_metadata = json.loads((run_dir/'circuit_metadata.json').read_text()) if (run_dir/'circuit_metadata.json').exists() else {}
    metadata = dict(planned_circuits=planned,completed_circuits=names,
                    failed_circuits=[r['circuit'] for r in failures],
                    pending_circuits=[n for n in planned if n not in names and n not in [r['circuit'] for r in failures]],
                    fault_model='TDF',method='CORE',low_power=True,threshold_percent=20,
                    dominance_reuse=False,core_extra_verify=False,performance_diagnostics=False,
                    source_commit=(run_dir/'source_commit.txt').read_text().strip(),
                    binary_sha256=(run_dir/'binary.sha256').read_text().split()[0],
                    raw_directory=str(run_dir),cpu_affinity=(run_dir/'affinity.txt').read_text().strip(),
                    wall_limit_seconds=int((run_dir/'wall_limit_seconds.txt').read_text()) if (run_dir/'wall_limit_seconds.txt').exists() else None,
                    circuit_metadata=source_metadata,
                    completed_timing='C clock() CPU and CLOCK_MONOTONIC wall; reported to 0.001 seconds',
                    timeout_timing='Process user+system CPU; active s5378 run sampled via procfs, s9234 via Bash waited child time',
                    small_warmups=1,small_repeats=7,medium_warmups=0,medium_repeats=1,
                    fdp='Pr[detection AND excitation AND power], denominator=2^(PI+initial DFF state)')
    snapshot = dict(metadata=metadata,summary=summary,measurements=raw,failures=failures)
    (run_dir/'summary.json').write_text(json.dumps(snapshot,indent=2,ensure_ascii=False)+'\n')
    with (run_dir/'summary.csv').open('w',newline='') as fp:
        writer=csv.DictWriter(fp,fieldnames=list(summary[0]) if summary else ['circuit']);writer.writeheader();writer.writerows(summary)
    # One graph-friendly CSV row per completed circuit or failed attempt.
    outcomes=[dict(r,status='completed') for r in summary]
    for f in failures:
        r={k:None for k in summary[0]} if summary else {}
        r.update(source_metadata.get(f['circuit'],{}))
        r.update(circuit=f['circuit'],status=f['status'],limit=30 if f['circuit'] in ('s5378','s9234') else 0,
                 threshold_percent=20,repeats=1,statistic=f['status'],cpu_s=f['cpu_s'],wall_s=f['wall_s'],
                 cpu_min_s=f['cpu_s'],cpu_max_s=f['cpu_s'])
        outcomes.append(r)
    columns=list(dict.fromkeys(k for r in outcomes for k in r))
    with (run_dir/'outcomes.csv').open('w',newline='') as fp:
        writer=csv.DictWriter(fp,fieldnames=columns);writer.writeheader();writer.writerows(outcomes)
    args.snapshot.parent.mkdir(parents=True,exist_ok=True)
    args.snapshot.write_text(json.dumps(snapshot,indent=2,ensure_ascii=False)+'\n')

    wb=Workbook();wb.remove(wb.active)
    wb.properties.title='TDF CORE low-power ON (20%)'
    formats=dict(text='General',int='#,##0',seconds='0.000',ratio='0.00%',avg='0.0000',prob='0.0000000000E+00')
    def sheet(name,columns,rows):
        ws=wb.create_sheet(name);ws.append([c[0] for c in columns])
        for row in rows:
            assert len(row)==len(columns),(name,row)
            ws.append(row)
        ws.freeze_panes='B2';ws.row_dimensions[1].height=36
        for i,(label,kind) in enumerate(columns,1):
            ws.column_dimensions[get_column_letter(i)].width=min(44,max(16,len(label)*1.6))
            cell=ws.cell(1,i);cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='24476B');cell.alignment=Alignment(wrap_text=True)
            for j in range(2,ws.max_row+1):
                cell=ws.cell(j,i);cell.number_format=formats[kind]
                if cell.value is not None and kind!='text':assert isinstance(cell.value,(int,float))
        table=Table(displayName=f'TDF{len(wb.worksheets)}',ref=f'A1:{get_column_letter(len(columns))}{ws.max_row}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    time_rows=[[r[k] for k in ('circuit','limit','repeats','statistic','cpu_s','cpu_min_s','cpu_max_s','wall_s')] for r in summary]
    time_rows += [[r['circuit'],30 if r['circuit'] in ('s5378','s9234') else 0,1,r['status'],
                  r['cpu_s'],r['cpu_s'],r['cpu_s'],r['wall_s']] for r in failures]
    sheet('実行時間',[('回路','text'),('limit（0は無制限）','int'),('測定回数','int'),('統計','text'),
        ('CORE CPU秒','seconds'),('CORE 最小CPU秒','seconds'),('CORE 最大CPU秒','seconds'),('経過秒','seconds')],
        time_rows)
    sheet('キューブ数',[('回路','text'),('limit（0は無制限）','int'),('CORE キューブ数','int')],
        [[r['circuit'],r['limit'],r['cubes']] for r in summary]+[[r['circuit'],30,None] for r in failures])
    sheet('故障完了',[('回路','text'),('全代表故障数','int'),('完了数','int'),('未完了数','int'),('完了率','ratio'),
        ('FDPゼロ数','int'),('正のFDPを得た数','int')],
        [[r['circuit'],r['representatives'],r['complete'],r['incomplete'],r['complete']/r['representatives'],
          r['zero_fdp'],r['positive_fdp']] for r in summary]+
        [[r['circuit'],source_metadata.get(r['circuit'],{}).get('representatives')]+[None]*5 for r in failures])
    sheet('X率',[('回路','text'),('入力ビット数','int'),('core抽出直後 X率','ratio'),('極小化後 X率','ratio'),
        ('削除判定対象／全入力','ratio'),('判定対象のうちX化','ratio')],
        [[r['circuit'],r['inputs']]+[r[k]/100 if r[k] is not None else None
          for k in ('core_x_pct','min_x_pct','tested_input_pct','deleted_core_pct')] for r in summary]+
        [[r['circuit'],source_metadata.get(r['circuit'],{}).get('inputs')]+[None]*4 for r in failures])
    sheet('判定回数',[('回路','text'),('生成キューブ数','int'),('coreケア平均／キューブ','avg'),('最終ケア平均／キューブ','avg'),
        ('否定側solve平均／キューブ','avg'),('全入力ビット位置数','int'),('coreケア総数','int'),('最終ケア総数','int'),
        ('削除試行総数','int'),('否定側solve総数','int')],
        [[r['circuit'],r['cubes']]+[r[k]/r['cubes'] if r['cubes'] else None
          for k in ('care_after_core','care_after_min','negative_solves')]+[r[k] for k in
          ('input_bits','care_after_core','care_after_min','care_after_core','negative_solves')] for r in summary]+
        [[r['circuit']]+[None]*9 for r in failures])
    sheet('時間内訳',[('回路','text'),('CORE DC秒','seconds'),('CORE 生成SAT秒','seconds'),('CORE BDD秒','seconds')],
        [[r[k] for k in ('circuit','dc_s','sat_s','bdd_s')] for r in summary]+[[r['circuit']]+[None]*3 for r in failures])
    sheet('電力条件',[('回路','text'),('閾値％','int'),('元信号線数','int'),('最大遷移数','int'),('自由入力数','int')],
        [[r[k] for k in ('circuit','threshold_percent','signals','budget','inputs')] for r in summary]+
        [[r['circuit'],20]+[source_metadata.get(r['circuit'],{}).get(k) for k in ('signals','budget','inputs')] for r in failures])
    if raw:
        keys=list(raw[0]);sheet('測定明細',[(k,'text' if k=='circuit' else 'seconds' if k.endswith('_s') else
            'avg' if k.endswith('_pct') else 'int') for k in keys],[[r[k] for k in keys] for r in raw])
    sheet('故障別FDP',[('回路','text'),('故障箇所','text'),('故障種別','text'),('キューブ数','int'),('complete','int'),
        ('FDP','prob'),('行種別','text')],fault_rows)
    notes=[['完走した回路',', '.join(names) or 'なし'],['未完走の回路',', '.join(metadata['failed_circuits']) or 'なし'],
        ['実行待ち・実行中の回路',', '.join(metadata['pending_circuits']) or 'なし'],
        ['故障モデル','遷移故障（TDF/LOC）。PIは2時刻共有、観測は2時刻目PPOのみ。'],
        ['比較対象','以前の paper_core_comparison.xlsx は縮退故障（SAF）・電力制約なし。故障モデルを区別する。'],
        ['方式','CORE。流用off、追加検証off。coreの再確認・極小化は通常通り実行。'],
        ['電力制約','正常回路の元信号線（分岐を含む）の値の変化数≤floor(信号線数×20/100)。'],
        ['FDP','全自由入力を分母にした検出∧励起∧電力条件の確率。電力条件を満たす入力だけへの正規化はしない。'],
        ['小規模','s27/s208: 完全列挙。ウォームアップ1回を除外して7回測定、時間は中央値。'],
        ['中規模','s5378/s9234: limit30、1回測定。complete=0のFDPは下界。'],
        ['タイムアウト','統計=timeout の時間は停止までの消費時間。完走CPU時間ではなく、その下限。故障完了・X率などの全件集計には使わない。'],
        ['保存済み行','未完走実行の保存済み代表故障数はCSVに書き出された行だけ。バッファに残った行は含まれず、実際に処理した数とは限らない。'],
        ['CPU最小・最大','中規模は1回なので同じ観測値。ばらつきを評価した値ではない。'],
        ['X率','全生成キューブの入力ビット位置を母数にする。キューブ0では未定義なので空欄。'],
        ['core直後','極小化を含む実行の中間統計。core単独で列挙した別実験ではない。'],
        ['FDPゼロ','complete=0でFDP=0の行があれば未発見を表し、テスト不能とは断定できない。'],
        ['精度','本体ログのCPU・経過時間は0.001秒刻み。時間内訳は一部前後処理を含まない。'],
        ['旧表とのCPU測定範囲','旧小規模SAFは子プロセスのuser+system全体（6桁）。今回の完走CPUは本体内部clock()（3桁）。旧中規模SAFは今回の完走時と同じ本体ログ。'],
        ['グラフ','グラフなし。数値セルとExcelテーブルのみ。'],['生データ',str(run_dir)],
        ['ソースコミット',metadata['source_commit']],['実行バイナリSHA256',metadata['binary_sha256']]]
    sheet('条件・注記',[('項目','text'),('内容','text')],notes)
    if failures:
        sheet('未完走実行',[('回路','text'),('run（0は準備実行）','int'),('状態','text'),('終了コード','int'),
            ('経過秒','seconds'),('消費CPU秒','seconds'),('保存済み代表故障数','int'),('保存済み完了数','int'),('保存済み未完了数','int')],
            [[r[k] for k in ('circuit','run','status','exit_code','wall_s','cpu_s','stored_representatives','stored_complete','stored_incomplete')] for r in failures])
    if partial_fault_rows:
        sheet('未完走の保存済みFDP',[('回路','text'),('故障箇所','text'),('故障種別','text'),('キューブ数','int'),
            ('complete','int'),('FDP','prob')],partial_fault_rows)
    args.output.parent.mkdir(parents=True,exist_ok=True);wb.save(args.output)
    reopened=load_workbook(args.output,data_only=True)
    assert reopened['実行時間'].max_row==len(summary)+len(failures)+1
    assert reopened['故障別FDP'].max_row==len(fault_rows)+1
    assert all(not ws._charts for ws in reopened)
    print(json.dumps(summary,ensure_ascii=False,indent=2))
    print('Excel:',args.output.resolve())


if __name__=='__main__':main()
